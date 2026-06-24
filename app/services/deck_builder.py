"""Render the deck-spec JSON returned by Claude into a .pptx file."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# 16:9 deck
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Palette — sell-side study-deck look.
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
TEXT = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x60, 0x60, 0x60)
BG = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD8, 0xDD, 0xE4)
LESSON_BG = RGBColor(0xE8, 0xF3, 0xE8)
LESSON_BORDER = RGBColor(0x2E, 0x7D, 0x32)
MISTAKE_BG = RGBColor(0xFB, 0xE8, 0xE6)
MISTAKE_BORDER = RGBColor(0xB7, 0x1C, 0x1C)
CHECK_BG = RGBColor(0xE6, 0xF0, 0xFA)
CHECK_BORDER = RGBColor(0x1F, 0x4E, 0x79)
TABLE_HEADER_BG = RGBColor(0x0B, 0x1F, 0x3A)
TABLE_ROW_ALT = RGBColor(0xF4, 0xF6, 0xFA)


# ─── helpers ─────────────────────────────────────────────────────────────────
def _set_text(frame, text: str, size: int, bold: bool = False,
              color: RGBColor = TEXT, align=PP_ALIGN.LEFT) -> None:
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _add_text_box(slide, left, top, width, height, text: str, size: int,
                  bold: bool = False, color: RGBColor = TEXT,
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    _set_text(tf, text, size, bold=bold, color=color, align=align)
    return box


def _add_bullets(slide, left, top, width, height, bullets: list[str],
                 size: int = 14) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"•  {bullet}"
        run.font.size = Pt(size)
        run.font.color.rgb = TEXT
        run.font.name = "Calibri"
        p.space_after = Pt(4)


def _add_title_bar(slide, section_title: str, slide_number: int,
                   total: int, title: str, subtitle: str | None) -> None:
    # Top color bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.35)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _set_text(
        bar.text_frame,
        f"  {section_title.upper()}",
        size=11,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Slide counter (right-aligned)
    _add_text_box(
        slide,
        SLIDE_W - Inches(1.5),
        Inches(0.05),
        Inches(1.4),
        Inches(0.25),
        f"Slide {slide_number} / {total}",
        size=10,
        color=RGBColor(0xCF, 0xD6, 0xE0),
        align=PP_ALIGN.RIGHT,
    )

    # Title
    _add_text_box(
        slide,
        Inches(0.5),
        Inches(0.5),
        SLIDE_W - Inches(1.0),
        Inches(0.7),
        title,
        size=28,
        bold=True,
        color=NAVY,
    )

    # Subtitle / horizontal rule
    if subtitle:
        _add_text_box(
            slide,
            Inches(0.5),
            Inches(1.15),
            SLIDE_W - Inches(1.0),
            Inches(0.4),
            subtitle,
            size=14,
            color=MUTED,
        )
        rule_top = Inches(1.6)
    else:
        rule_top = Inches(1.25)

    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), rule_top,
        SLIDE_W - Inches(1.0), Emu(12700),  # ~1pt
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = RULE
    rule.line.fill.background()


def _add_callout(slide, left, top, width, height, label: str, text: str,
                 bg: RGBColor, border: RGBColor) -> None:
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.color.rgb = border
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)

    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = label
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = border
    r1.font.name = "Calibri"

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = text
    r2.font.size = Pt(11)
    r2.font.color.rgb = TEXT
    r2.font.name = "Calibri"
    p2.space_before = Pt(2)


def _add_formula_block(slide, left, top, width, height, formula: dict) -> None:
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF7, 0xF8, 0xFB)
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)

    name = formula.get("name", "Formula")
    expression = formula.get("expression", "")
    inputs = formula.get("inputs", "")
    result = formula.get("result", "")
    interp = formula.get("interpretation", "")

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = f"FORMULA — {name}"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.name = "Calibri"

    def _row(label: str, value: str, bold: bool = False):
        if not value:
            return
        para = tf.add_paragraph()
        para.space_before = Pt(3)
        rl = para.add_run()
        rl.text = f"{label}  "
        rl.font.size = Pt(11)
        rl.font.color.rgb = MUTED
        rl.font.bold = True
        rl.font.name = "Calibri"
        rv = para.add_run()
        rv.text = value
        rv.font.size = Pt(12 if bold else 11)
        rv.font.bold = bold
        rv.font.color.rgb = TEXT
        rv.font.name = "Consolas"

    _row("Expression:", expression, bold=True)
    _row("Inputs:", inputs)
    _row("Result:", result, bold=True)
    if interp:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        ri = p2.add_run()
        ri.text = interp
        ri.font.size = Pt(11)
        ri.font.italic = True
        ri.font.color.rgb = TEXT
        ri.font.name = "Calibri"


def _add_table(slide, left, top, width, height, table_spec: dict) -> None:
    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    caption = table_spec.get("caption")

    if not headers or not rows:
        return

    cols = len(headers)
    n_rows = len(rows) + 1  # + header

    # Constrain rows shown to keep the slide readable
    if n_rows > 16:
        rows = rows[:15]
        n_rows = 16

    table_shape = slide.shapes.add_table(n_rows, cols, left, top, width, height)
    table = table_shape.table

    # Header row
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.text = ""
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = str(h)
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = "Calibri"

    # Body
    for r_idx, row in enumerate(rows, start=1):
        for c in range(cols):
            cell = table.cell(r_idx, c)
            if r_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ROW_ALT
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            val = row[c] if c < len(row) else ""
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(10)
            run.font.color.rgb = TEXT
            run.font.name = "Calibri"

    if caption:
        _add_text_box(
            slide,
            left,
            top + height + Inches(0.05),
            width,
            Inches(0.3),
            caption,
            size=10,
            color=MUTED,
        )


# ─── slide builders ──────────────────────────────────────────────────────────
def _blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(layout)
    # Force white background.
    bg_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg_rect.fill.solid()
    bg_rect.fill.fore_color.rgb = BG
    bg_rect.line.fill.background()
    # Send to back by removing/re-adding ordering not exposed by python-pptx;
    # in practice every other shape we add lands above it.
    return slide


def _build_cover(prs: Presentation, spec: dict) -> None:
    slide = _blank_slide(prs)
    company = spec.get("company", {})
    name = company.get("name", "Company")
    ticker = company.get("ticker", "—")
    sector = company.get("sector", "")
    industry = company.get("industry", "")
    as_of = company.get("as_of", "")
    thesis_tag = spec.get("thesis_tag", "")
    confidence = spec.get("confidence_pct", "")
    one_line = spec.get("one_line_thesis", "")

    # Big color band on the left
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), SLIDE_H
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()

    _add_text_box(
        slide, Inches(0.9), Inches(0.8), SLIDE_W - Inches(1.4), Inches(0.45),
        "STUDY DECK — FUNDAMENTAL ANALYSIS",
        size=14, bold=True, color=ACCENT,
    )
    _add_text_box(
        slide, Inches(0.9), Inches(1.4), SLIDE_W - Inches(1.4), Inches(1.4),
        name,
        size=48, bold=True, color=NAVY,
    )
    _add_text_box(
        slide, Inches(0.9), Inches(2.7), SLIDE_W - Inches(1.4), Inches(0.5),
        f"{ticker}  ·  {sector}  ·  {industry}",
        size=18, color=MUTED,
    )
    _add_text_box(
        slide, Inches(0.9), Inches(3.4), SLIDE_W - Inches(1.4), Inches(0.4),
        f"As of {as_of}",
        size=13, color=MUTED,
    )

    # Thesis teaser box
    if one_line:
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.9), Inches(4.4),
            SLIDE_W - Inches(1.4), Inches(1.6),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF7, 0xF8, 0xFB)
        box.line.color.rgb = ACCENT
        box.line.width = Pt(1)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = f"THESIS  ·  {thesis_tag}  ·  Confidence {confidence}%"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        r.font.name = "Calibri"
        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        r2 = p2.add_run()
        r2.text = one_line
        r2.font.size = Pt(16)
        r2.font.color.rgb = TEXT
        r2.font.name = "Calibri"

    _add_text_box(
        slide, Inches(0.9), SLIDE_H - Inches(0.6),
        SLIDE_W - Inches(1.4), Inches(0.4),
        "Generated by the Equity Deck Agent — for educational use; not investment advice.",
        size=10, color=MUTED,
    )


def _build_content_slide(prs: Presentation, slide_spec: dict,
                         total: int) -> None:
    slide = _blank_slide(prs)
    _add_title_bar(
        slide,
        slide_spec.get("section_title", ""),
        slide_spec.get("slide_number", 0),
        total,
        slide_spec.get("title", ""),
        slide_spec.get("subtitle"),
    )

    content_top = Inches(1.8)
    content_left = Inches(0.5)
    content_width = SLIDE_W - Inches(1.0)
    cursor = content_top

    body = slide_spec.get("body")
    if body:
        _add_text_box(
            slide, content_left, cursor, content_width, Inches(1.0),
            body, size=13, color=TEXT,
        )
        cursor += Inches(0.9)

    bullets = slide_spec.get("bullets") or []
    if bullets:
        height = Inches(0.35) * len(bullets) + Inches(0.15)
        _add_bullets(slide, content_left, cursor, content_width, height,
                     bullets, size=13)
        cursor += height + Inches(0.1)

    formula = slide_spec.get("formula")
    if formula:
        _add_formula_block(
            slide, content_left, cursor, content_width, Inches(1.5), formula,
        )
        cursor += Inches(1.6)

    table = slide_spec.get("table")
    if table:
        available = SLIDE_H - cursor - Inches(1.5)
        height = max(Inches(1.5), min(Inches(3.5), available))
        _add_table(slide, content_left, cursor, content_width, height, table)
        cursor += height + Inches(0.2)

    # Pedagogical callouts at the bottom, side by side.
    callouts = []
    if slide_spec.get("lesson_card"):
        callouts.append(("LESSON CARD", slide_spec["lesson_card"],
                         LESSON_BG, LESSON_BORDER))
    if slide_spec.get("common_mistake"):
        callouts.append(("COMMON MISTAKE", slide_spec["common_mistake"],
                         MISTAKE_BG, MISTAKE_BORDER))
    if slide_spec.get("quick_check"):
        callouts.append(("QUICK CHECK", slide_spec["quick_check"],
                         CHECK_BG, CHECK_BORDER))

    if callouts:
        n = len(callouts)
        gap = Inches(0.15)
        width = (content_width - gap * (n - 1)) / n
        top = SLIDE_H - Inches(1.3)
        for i, (label, text, bg, border) in enumerate(callouts):
            left = content_left + (width + gap) * i
            _add_callout(slide, left, top, width, Inches(1.1),
                         label, text, bg, border)


def _build_how_to_read(prs: Presentation, slide_spec: dict, total: int) -> None:
    """Render the icons legend as three callouts so readers see them upfront."""
    slide = _blank_slide(prs)
    _add_title_bar(
        slide,
        slide_spec.get("section_title", ""),
        slide_spec.get("slide_number", 0),
        total,
        slide_spec.get("title", "How to read this deck"),
        slide_spec.get("subtitle"),
    )
    bullets = slide_spec.get("bullets") or []
    if bullets:
        _add_bullets(
            slide, Inches(0.5), Inches(1.8),
            SLIDE_W - Inches(1.0), Inches(3.2), bullets, size=14,
        )

    # Legend row
    legend_top = SLIDE_H - Inches(2.0)
    legend_width = (SLIDE_W - Inches(1.0) - Inches(0.4)) / 3
    callouts = [
        ("LESSON CARD",
         "Transferable principle — what this case teaches about fundamental analysis broadly.",
         LESSON_BG, LESSON_BORDER),
        ("COMMON MISTAKE",
         "A specific cognitive or methodological trap the metric on this slide invites.",
         MISTAKE_BG, MISTAKE_BORDER),
        ("QUICK CHECK",
         "A question you should be able to answer from this slide alone.",
         CHECK_BG, CHECK_BORDER),
    ]
    for i, (label, text, bg, border) in enumerate(callouts):
        left = Inches(0.5) + (legend_width + Inches(0.2)) * i
        _add_callout(slide, left, legend_top, legend_width, Inches(1.6),
                     label, text, bg, border)


def _build_quiz_answers(prs: Presentation, slide_spec: dict, total: int) -> None:
    slide = _blank_slide(prs)
    _add_title_bar(
        slide,
        slide_spec.get("section_title", ""),
        slide_spec.get("slide_number", 0),
        total,
        slide_spec.get("title", "Answer Key"),
        slide_spec.get("subtitle"),
    )
    bullets = slide_spec.get("bullets") or []
    if bullets:
        _add_bullets(
            slide, Inches(0.5), Inches(1.8),
            SLIDE_W - Inches(1.0), Inches(5.0), bullets, size=13,
        )


# ─── entry point ─────────────────────────────────────────────────────────────
def build_pptx(spec: dict[str, Any], output_path: Path | str) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slides = spec.get("slides", [])
    total = len(slides)

    for slide_spec in slides:
        kind = slide_spec.get("kind", "content")
        if kind == "cover":
            _build_cover(prs, spec)
        elif kind == "how_to_read":
            _build_how_to_read(prs, slide_spec, total)
        elif kind == "quiz_answers":
            _build_quiz_answers(prs, slide_spec, total)
        else:
            _build_content_slide(prs, slide_spec, total)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
